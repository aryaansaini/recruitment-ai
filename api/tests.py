import io
import zipfile
from unittest import mock

from django.contrib.auth.models import User
from django.core.cache import cache
from django.db import IntegrityError, transaction
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APITestCase

from google.adk.models import BaseLlm
from google.adk.models.llm_response import LlmResponse
from google.genai import types
from pydantic import ValidationError

from .adk_agents import SkillProfile, extract_skills
from .llm import LLMError, MatchResult, ScreeningResult, generate_structured
from .models import Application, Candidate, Job, ResumeChunk
from .resume_parser import (
    ResumeParseError, UnsupportedResumeType, chunk_resume, extract_email, extract_phone,
    extract_text, parse_skill_list, split_sections,
)
from .services import rebuild_resume_chunks


def make_candidate(**kw):
    data = dict(name='Rahul', email='rahul@example.com', phone='9876543210',
                skills='Python, Django, SQL', experience=3,
                resume_text='Backend dev with Django and SQL experience')
    data.update(kw)
    return Candidate.objects.create(**data)


def make_job(**kw):
    data = dict(title='Backend Dev', required_skills='Python, Django',
                min_experience=2, description='Build APIs')
    data.update(kw)
    return Job.objects.create(**data)


class BaseAPITest(APITestCase):
    def setUp(self):
        cache.clear()  # reset throttle counters between tests
        self.user = User.objects.create_user('recruiter', password='pass12345')

    def login(self):
        self.client.force_authenticate(self.user)


class AuthTests(BaseAPITest):
    def test_login_returns_tokens(self):
        r = self.client.post('/api/auth/login/', {'username': 'recruiter', 'password': 'pass12345'}, format='json')
        self.assertEqual(r.status_code, 200)
        self.assertIn('access', r.data)
        self.assertIn('refresh', r.data)

    def test_login_wrong_password_is_401(self):
        r = self.client.post('/api/auth/login/', {'username': 'recruiter', 'password': 'wrong'}, format='json')
        self.assertEqual(r.status_code, 401)

    def test_token_grants_access(self):
        r = self.client.post('/api/auth/login/', {'username': 'recruiter', 'password': 'pass12345'}, format='json')
        self.client.credentials(HTTP_AUTHORIZATION='Bearer ' + r.data['access'])
        self.assertEqual(self.client.get('/api/candidates/').status_code, 200)

    def test_garbage_token_is_401(self):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer not-a-real-token')
        self.assertEqual(self.client.get('/api/candidates/').status_code, 401)

    def test_login_is_rate_limited(self):
        codes = [self.client.post('/api/auth/login/', {'username': 'recruiter', 'password': 'bad'},
                                  format='json').status_code for _ in range(12)]
        self.assertIn(429, codes)


class ProtectionTests(BaseAPITest):
    def test_all_data_endpoints_require_login(self):
        for url in ['/api/candidates/', '/api/jobs/', '/api/logs/']:
            self.assertEqual(self.client.get(url).status_code, 401, url)
        for url in ['/api/add-job/', '/api/add-candidate/', '/api/run-pipeline/', '/api/parse-resume/']:
            self.assertEqual(self.client.post(url, {}, format='json').status_code, 401, url)

    def test_homepage_is_public(self):
        self.assertEqual(self.client.get('/').status_code, 200)


class CreateTests(BaseAPITest):
    def setUp(self):
        super().setUp()
        self.login()

    def test_add_job_success(self):
        r = self.client.post('/api/add-job/', {
            'title': 'Dev', 'required_skills': 'Python', 'min_experience': '2', 'description': 'x'}, format='json')
        self.assertEqual(r.status_code, 201)
        self.assertIn('id', r.data)

    def test_add_job_rejects_negative_experience_and_missing_title(self):
        r = self.client.post('/api/add-job/', {
            'required_skills': 'Python', 'min_experience': -1, 'description': 'x'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('title', r.data)
        self.assertIn('min_experience', r.data)

    def test_add_candidate_success(self):
        r = self.client.post('/api/add-candidate/', {
            'name': 'A', 'email': 'a@b.com', 'phone': '123', 'skills': 'Python',
            'experience': '3', 'resume_text': 'resume'}, format='json')
        self.assertEqual(r.status_code, 201)
        self.assertEqual(Candidate.objects.count(), 1)

    def test_add_candidate_rejects_bad_email_and_empty_resume(self):
        r = self.client.post('/api/add-candidate/', {
            'name': 'A', 'email': 'not-an-email', 'phone': '123', 'skills': 'Python',
            'experience': '3', 'resume_text': ''}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('email', r.data)
        self.assertIn('resume_text', r.data)
        self.assertEqual(Candidate.objects.count(), 0)


def ai_ok(score=85, match=True):
    """Fake generate_structured: returns a valid AI answer for each schema."""
    def fake(prompt, schema, **kw):
        if schema is ScreeningResult:
            return ScreeningResult(score=score, reason='Strong Django and SQL evidence')
        return MatchResult(is_match=match, match_percent=90 if match else 20, reason='Skills overlap')
    return fake


class PipelineTests(BaseAPITest):
    def setUp(self):
        super().setUp()
        self.login()

    def run_pipeline(self, c, j):
        return self.client.post('/api/run-pipeline/', {'candidate_id': c.id, 'job_id': j.id}, format='json')

    def patch_ai(self, **kw):
        p = mock.patch('api.agents.generate_structured', **kw)
        p.start()
        self.addCleanup(p.stop)

    def test_unknown_ids_return_404(self):
        r = self.client.post('/api/run-pipeline/', {'candidate_id': 999, 'job_id': 999}, format='json')
        self.assertEqual(r.status_code, 404)

    def test_missing_fields_return_400(self):
        r = self.client.post('/api/run-pipeline/', {'candidate_id': 1}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('job_id', r.data)

    def test_ai_success_stores_reason_and_no_fallback_flag(self):
        self.patch_ai(side_effect=ai_ok(score=85))
        c, j = make_candidate(), make_job()
        r = self.run_pipeline(c, j)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['score'], 85)
        self.assertFalse(r.data['fallback_used'])
        self.assertIn('Strong Django', r.data['reason'])
        app = Application.objects.get(candidate=c, job=j)
        self.assertFalse(app.fallback_used)
        self.assertIn('Strong Django', app.reason)
        self.assertEqual(app.status, 'interview_scheduled')

    def test_ai_failure_uses_fallback_and_is_flagged(self):
        self.patch_ai(side_effect=LLMError('quota exceeded'))
        c, j = make_candidate(), make_job()
        r = self.run_pipeline(c, j)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data['fallback_used'])
        self.assertIn('AI unavailable', r.data['reason'])
        self.assertTrue(Application.objects.get(candidate=c, job=j).fallback_used)

    def test_fallback_score_never_exceeds_100(self):
        self.patch_ai(side_effect=LLMError('down'))
        c = make_candidate(skills=','.join(f's{i}' for i in range(30)), experience=40)
        r = self.run_pipeline(c, make_job())
        self.assertLessEqual(r.data['score'], 100)

    def test_ai_low_score_rejects_without_interview(self):
        self.patch_ai(side_effect=ai_ok(score=20))
        c, j = make_candidate(), make_job()
        r = self.run_pipeline(c, j)
        self.assertEqual(r.data['status'], 'rejected')
        self.assertIsNone(r.data['interview_slot'])

    def test_ai_no_match_rejects(self):
        self.patch_ai(side_effect=ai_ok(score=80, match=False))
        c, j = make_candidate(), make_job()
        r = self.run_pipeline(c, j)
        self.assertEqual(r.data['status'], 'rejected')

    def test_weak_candidate_rejected_by_fallback(self):
        self.patch_ai(side_effect=LLMError('down'))
        c = make_candidate(skills='Excel', experience=0, resume_text='data entry')
        r = self.run_pipeline(c, make_job())
        self.assertEqual(r.data['status'], 'rejected')
        self.assertIsNone(r.data['interview_slot'])

    def test_pipeline_is_rate_limited(self):
        self.patch_ai(side_effect=LLMError('down'))
        c, j = make_candidate(), make_job()
        codes = [self.run_pipeline(c, j).status_code for _ in range(12)]
        self.assertIn(429, codes)


class LLMWrapperTests(BaseAPITest):
    def _client_returning(self, **resp_attrs):
        client = mock.Mock()
        client.models.generate_content.return_value = mock.Mock(**resp_attrs)
        return client

    def test_returns_parsed_object(self):
        parsed = ScreeningResult(score=70, reason='ok')
        with mock.patch('api.llm._get_client', return_value=self._client_returning(parsed=parsed, text='')):
            self.assertEqual(generate_structured('p', ScreeningResult).score, 70)

    def test_falls_back_to_json_text(self):
        with mock.patch('api.llm._get_client',
                        return_value=self._client_returning(parsed=None, text='{"score": 55, "reason": "fine"}')):
            self.assertEqual(generate_structured('p', ScreeningResult).score, 55)

    def test_score_is_clamped_to_0_100(self):
        self.assertEqual(ScreeningResult(score=250, reason='x').score, 100)
        self.assertEqual(ScreeningResult(score=-5, reason='x').score, 0)

    def test_invalid_json_raises_llmerror(self):
        with mock.patch('api.llm._get_client', return_value=self._client_returning(parsed=None, text='not json')):
            with self.assertRaises(LLMError):
                generate_structured('p', ScreeningResult)

    def test_api_exception_raises_llmerror(self):
        client = mock.Mock()
        client.models.generate_content.side_effect = RuntimeError('503 unavailable')
        with mock.patch('api.llm._get_client', return_value=client):
            with self.assertRaises(LLMError):
                generate_structured('p', ScreeningResult)

    def test_missing_api_key_raises_llmerror(self):
        import api.llm as llm
        with mock.patch.object(llm, '_client', None), self.settings(GEMINI_API_KEY=''):
            with self.assertRaises(LLMError):
                generate_structured('p', ScreeningResult)


class StubLlm(BaseLlm):
    """Offline stand-in for Gemini so ADK can be tested without network."""
    model: str = 'stub'
    reply: str = ''

    @classmethod
    def supported_models(cls):
        return ['stub']

    async def generate_content_async(self, llm_request, stream=False):
        yield LlmResponse(content=types.Content(role='model', parts=[types.Part(text=self.reply)]))


class ADKTests(BaseAPITest):
    def test_adk_agent_runs_from_sync_django_and_returns_structured_output(self):
        stub = StubLlm(reply='{"skills": ["python", "django"], "years_experience": 3, "summary": "Backend dev"}')
        profile = extract_skills('Rahul, 3 years Python and Django', model=stub)
        self.assertIsInstance(profile, SkillProfile)
        self.assertEqual(profile.skills, ['python', 'django'])
        self.assertEqual(profile.years_experience, 3)

    def test_adk_bad_model_output_is_rejected(self):
        stub = StubLlm(reply='{"skills": "not-a-list"}')
        with self.assertRaises(ValidationError):
            extract_skills('resume', model=stub)


SAMPLE_RESUME = """Rahul Sharma
rahul.sharma@example.com | +91 98765 43210

PROFESSIONAL SUMMARY
Backend developer with 3 years of experience building REST APIs.

TECHNICAL SKILLS
Languages: Python, Java
Frameworks: Django, DRF | SQL, Docker

WORK EXPERIENCE
Backend Developer, Acme Corp (2022 - 2025)
- Built REST APIs with Django and PostgreSQL
- Reduced API latency by 40 percent

PROJECTS
Recruitment AI: multi-agent hiring pipeline using Gemini and ADK.

EDUCATION
B.Tech Computer Science, 2022
"""


def make_docx(text, table_rows=None):
    from docx import Document
    doc = Document()
    for line in text.split('\n'):
        doc.add_paragraph(line)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(lines):
    """Build a tiny valid one-page PDF (no extra dependency needed in tests)."""
    esc = lambda t: t.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
    content = 'BT /F1 11 Tf 50 780 Td 14 TL ' + ' '.join(f'({esc(l)}) Tj T*' for l in lines) + ' ET'
    objs = [
        '<< /Type /Catalog /Pages 2 0 R >>',
        '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R '
        '/Resources << /Font << /F1 5 0 R >> >> >>',
        f'<< /Length {len(content)} >>\nstream\n{content}\nendstream',
        '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    out, offsets = b'%PDF-1.4\n', []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f'{i} 0 obj\n{obj}\nendobj\n'.encode('latin-1')
    xref = len(out)
    out += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode()
    for off in offsets:
        out += f'{off:010d} 00000 n \n'.encode()
    out += f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
    return out


class ExtractTextTests(BaseAPITest):
    def test_txt(self):
        self.assertIn('Rahul', extract_text(SAMPLE_RESUME.encode(), 'cv.txt'))

    def test_txt_with_non_utf8_encoding(self):
        self.assertIn('Zoë', extract_text('Zoë Developer'.encode('cp1252'), 'cv.txt'))

    def test_docx_paragraphs(self):
        text = extract_text(make_docx(SAMPLE_RESUME), 'cv.docx')
        self.assertIn('TECHNICAL SKILLS', text)
        self.assertIn('Django', text)

    def test_docx_includes_table_text(self):
        data = make_docx('Header line', table_rows=[['Skill', 'Python'], ['Skill', 'SQL']])
        text = extract_text(data, 'cv.docx')
        self.assertIn('Python', text)
        self.assertIn('SQL', text)

    def test_pdf(self):
        text = extract_text(make_pdf(['Rahul Sharma', 'SKILLS', 'Python, Django']), 'cv.pdf')
        self.assertIn('Rahul Sharma', text)
        self.assertIn('Python, Django', text)

    def test_unsupported_extension(self):
        with self.assertRaises(UnsupportedResumeType):
            extract_text(b'MZ', 'virus.exe')

    def test_fake_pdf_and_fake_docx_are_rejected_by_content_not_name(self):
        with self.assertRaises(ResumeParseError):
            extract_text(b'this is not a pdf', 'cv.pdf')
        with self.assertRaises(ResumeParseError):
            extract_text(b'this is not a docx', 'cv.docx')

    def test_empty_file_is_rejected(self):
        with self.assertRaises(ResumeParseError):
            extract_text(b'   \n  ', 'cv.txt')

    def test_zip_bomb_docx_is_rejected(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
            zf.writestr('word/document.xml', b'0' * (60 * 1024 * 1024))   # 60 MB unpacked, tiny packed
        self.assertLess(len(buf.getvalue()), 1024 * 1024)
        with self.assertRaises(ResumeParseError):
            extract_text(buf.getvalue(), 'bomb.docx')


class ContactAndSkillsTests(BaseAPITest):
    def test_email_and_phone(self):
        self.assertEqual(extract_email(SAMPLE_RESUME), 'rahul.sharma@example.com')
        self.assertEqual(extract_phone(SAMPLE_RESUME), '+919876543210')

    def test_no_contact_returns_empty(self):
        self.assertEqual(extract_email('no contact here'), '')
        self.assertEqual(extract_phone('born in 1999, 3 years exp'), '')

    def test_skill_list_handles_labels_separators_and_duplicates(self):
        got = parse_skill_list('Languages: Python, Java\n- Django | SQL\n• python\nFrameworks: DRF; Docker')
        self.assertEqual(got, ['Python', 'Java', 'Django', 'SQL', 'DRF', 'Docker'])

    def test_skill_list_drops_long_sentences(self):
        self.assertEqual(parse_skill_list('Python, I am a very hard working team player who loves code'), ['Python'])


class SplitSectionsTests(BaseAPITest):
    def test_standard_resume(self):
        sections = split_sections(SAMPLE_RESUME)
        self.assertEqual(list(sections), ['header', 'summary', 'skills', 'experience', 'projects', 'education'])
        self.assertIn('rahul.sharma@example.com', sections['header'])
        self.assertIn('Django', sections['skills'])
        self.assertIn('Reduced API latency', sections['experience'])

    def test_inline_heading(self):
        sections = split_sections('Asha\nSkills: Python, SQL\nExperience: 2 years at Foo')
        self.assertEqual(sections['skills'], 'Python, SQL')
        self.assertIn('2 years at Foo', sections['experience'])

    def test_heading_variants(self):
        sections = split_sections('Name\nWork Experience:\nDid things\nTech Stack\nPython')
        self.assertIn('Did things', sections['experience'])
        self.assertIn('Python', sections['skills'])

    def test_no_headings_everything_is_header(self):
        self.assertEqual(list(split_sections('just some plain text\nwith two lines')), ['header'])

    def test_heading_word_inside_sentence_is_not_a_heading(self):
        sections = split_sections('Name\nSKILLS\nI have strong skills and education in testing')
        self.assertEqual(list(sections), ['header', 'skills'])


class ChunkingTests(BaseAPITest):
    def test_header_is_never_chunked(self):
        chunks = chunk_resume(SAMPLE_RESUME)
        self.assertNotIn('header', {c.section for c in chunks})
        self.assertNotIn('rahul.sharma@example.com', ' '.join(c.text for c in chunks))

    def test_chunks_have_sections_and_sequential_indexes(self):
        chunks = chunk_resume(SAMPLE_RESUME)
        self.assertEqual([c.index for c in chunks], list(range(len(chunks))))
        self.assertEqual({c.section for c in chunks}, {'summary', 'skills', 'experience', 'projects', 'education'})

    def test_long_section_is_split_within_limit_without_losing_lines(self):
        bullets = [f'- Built feature number {i} using Django and PostgreSQL' for i in range(40)]
        text = 'Name\nEXPERIENCE\n' + '\n'.join(bullets)
        chunks = chunk_resume(text, max_chars=300, overlap=60)
        self.assertGreater(len(chunks), 3)
        self.assertTrue(all(len(c.text) <= 300 for c in chunks))
        joined = '\n'.join(c.text for c in chunks)
        for b in bullets:
            self.assertIn(b, joined)

    def test_consecutive_chunks_overlap(self):
        bullets = [f'- item {i:02d} ' + 'x' * 30 for i in range(30)]
        chunks = chunk_resume('Name\nSKILLS\n' + '\n'.join(bullets), max_chars=200, overlap=80)
        first_tail = chunks[0].text.split('\n')[-1]
        self.assertIn(first_tail, chunks[1].text)

    def test_single_giant_line_is_hard_wrapped(self):
        chunks = chunk_resume('Name\nPROJECTS\n' + 'word ' * 500, max_chars=200)
        self.assertGreater(len(chunks), 5)
        self.assertTrue(all(len(c.text) <= 200 for c in chunks))

    def test_unbroken_giant_token_is_split(self):
        chunks = chunk_resume('Name\nPROJECTS\n' + 'a' * 1000, max_chars=200)
        self.assertTrue(all(len(c.text) <= 200 for c in chunks))
        self.assertEqual(sum(len(c.text) for c in chunks), 1000)

    def test_empty_and_header_only_text(self):
        self.assertEqual(chunk_resume(''), [])
        self.assertEqual(chunk_resume('only a header line'), [])

    def test_embed_text_has_section_label(self):
        chunk = chunk_resume('Name\nSKILLS\nPython')[0]
        self.assertEqual(chunk.embed_text, 'Skills: Python')


class ParseResumeTests(BaseAPITest):
    def setUp(self):
        super().setUp()
        self.login()

    def upload(self, name, data, content_type='application/octet-stream'):
        return self.client.post('/api/parse-resume/', {'resume': SimpleUploadedFile(name, data, content_type=content_type)},
                                format='multipart')

    def test_no_file_is_400(self):
        self.assertEqual(self.client.post('/api/parse-resume/', {}, format='multipart').status_code, 400)

    def test_unsupported_type_is_400(self):
        self.assertEqual(self.upload('resume.exe', b'MZ...').status_code, 400)

    def test_oversized_file_is_400(self):
        self.assertEqual(self.upload('big.pdf', b'0' * (5 * 1024 * 1024 + 1)).status_code, 400)

    def test_corrupt_pdf_is_422(self):
        self.assertEqual(self.upload('bad.pdf', b'this is not a pdf').status_code, 422)

    def test_docx_upload_returns_structured_data(self):
        r = self.upload('cv.docx', make_docx(SAMPLE_RESUME))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['email'], 'rahul.sharma@example.com')
        self.assertEqual(r.data['phone'], '+919876543210')
        self.assertIn('Django', r.data['skills'])
        self.assertIn('experience', r.data['sections'])
        self.assertNotIn('header', r.data['sections'])
        self.assertGreater(r.data['chunk_count'], 3)
        self.assertIn('Rahul Sharma', r.data['text'])

    def test_pdf_upload_works(self):
        r = self.upload('cv.pdf', make_pdf(['Asha Rao', 'SKILLS', 'Python, SQL']))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['skills'], ['Python', 'SQL'])

    def test_txt_upload_works(self):
        r = self.upload('cv.txt', SAMPLE_RESUME.encode())
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['email'], 'rahul.sharma@example.com')


class ChunkStorageTests(BaseAPITest):
    def setUp(self):
        super().setUp()
        self.login()

    def test_add_candidate_stores_chunks(self):
        r = self.client.post('/api/add-candidate/', {
            'name': 'Rahul', 'email': 'r@x.com', 'phone': '123', 'skills': 'Python',
            'experience': 3, 'resume_text': SAMPLE_RESUME}, format='json')
        self.assertEqual(r.status_code, 201)
        c = Candidate.objects.get(id=r.data['id'])
        self.assertEqual(c.chunks.count(), r.data['chunks'])
        self.assertGreater(r.data['chunks'], 3)
        self.assertIn('skills', set(c.chunks.values_list('section', flat=True)))

    def test_rebuild_replaces_old_chunks_instead_of_duplicating(self):
        c = make_candidate(resume_text=SAMPLE_RESUME)
        first = rebuild_resume_chunks(c)
        second = rebuild_resume_chunks(c)
        self.assertEqual(first, second)
        self.assertEqual(c.chunks.count(), first)

    def test_deleting_candidate_deletes_chunks(self):
        c = make_candidate(resume_text=SAMPLE_RESUME)
        rebuild_resume_chunks(c)
        c.delete()
        self.assertEqual(ResumeChunk.objects.count(), 0)


class ModelTests(BaseAPITest):
    def test_one_application_per_candidate_job_pair(self):
        c, j = make_candidate(), make_job()
        Application.objects.create(candidate=c, job=j)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Application.objects.create(candidate=c, job=j)

    def test_candidate_can_apply_to_multiple_jobs(self):
        c = make_candidate()
        Application.objects.create(candidate=c, job=make_job(title='A'))
        Application.objects.create(candidate=c, job=make_job(title='B'))
        self.assertEqual(c.applications.count(), 2)